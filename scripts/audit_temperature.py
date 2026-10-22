"""Audit the 2024 quarter-hour temperature workbook without modifying it.

Run: python scripts/audit_temperature.py --workbook input.xlsx --output audit-results
Requires openpyxl and pandas. Output is audit-results/summary.json.
The expected schema is datetime/T_out and the reference grid is calendar 2024.
The Brussels localization check is hypothetical; it does not infer a timezone.
Author identities, saved paths, and annotation contents are excluded from output.
"""
import argparse
from collections import Counter
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import re
from zipfile import ZipFile

import openpyxl
import pandas as pd


def iso(value):
    return value.isoformat() if hasattr(value, "isoformat") else value


def inspect(path):
    digest_before = hashlib.sha256(path.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    result = {
        "source_file_name": path.name, "source_bytes": path.stat().st_size,
        "sha256": digest_before, "software": {
            "openpyxl": openpyxl.__version__, "pandas": pd.__version__},
        "excel_epoch": wb.epoch.isoformat(),
        "metadata_check": {
            "descriptive_properties_present": {
                name: bool(getattr(wb.properties, name, None))
                for name in ("title", "description", "subject", "keywords")
            },
            "property_values_and_saved_locations_omitted": True,
        },
        "sheet_names": wb.sheetnames, "defined_name_count": len(wb.defined_names),
        "sheets": [],
    }
    with ZipFile(path) as z:
        result["zip_members"] = z.namelist()
        result["embedded_external_data_connections"] = any(
            re.search(r"(^xl/connections|^xl/externalLinks|^xl/queryTables)", n)
            for n in z.namelist())
        result["custom_property_parts"] = [n for n in z.namelist() if "custom" in n.lower()]
    for ws in wb:
        all_rows = list(ws.iter_rows())
        header = [c.value for c in all_rows[0]]
        if header != ["datetime", "T_out"]:
            raise ValueError(f"Sheet {ws.title!r}: expected exactly datetime/T_out columns")
        rows = [[c.value for c in r] for r in all_rows[1:]]
        timestamps = [r[0] for r in rows]
        values = [r[1] for r in rows]
        valid_timestamps = [t for t in timestamps if isinstance(t, datetime)]
        valid_numeric = [v for v in values if isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)]
        if not valid_timestamps or not valid_numeric:
            raise ValueError(f"Sheet {ws.title!r}: requires datetime and finite numeric observations")
        idx = pd.DatetimeIndex(valid_timestamps)
        deltas = Counter((b-a).total_seconds() for a, b in zip(valid_timestamps, valid_timestamps[1:]))
        expected = pd.date_range("2024-01-01 00:00", "2024-12-31 23:45", freq="15min")
        per_day = Counter(t.date().isoformat() for t in valid_timestamps)
        # This is a hypothetical localization diagnostic, NOT a timezone inference.
        ambiguous = idx.tz_localize("Europe/Brussels", ambiguous="NaT", nonexistent="shift_forward").isna()
        nonexistent = idx.tz_localize("Europe/Brussels", ambiguous=True, nonexistent="NaT").isna()
        comments = []
        hyperlinks = []
        formulas = []
        styles = Counter()
        fills = Counter()
        data_types = Counter()
        for row in all_rows:
            for c in row:
                styles[f"{c.column_letter}:style_{c.style_id}"] += 1
                fills[f"pattern={c.fill.patternType},fg_type={c.fill.fgColor.type},fg={c.fill.fgColor.index}"] += 1
                data_types[f"{c.column_letter}:{c.data_type}"] += 1
                if c.comment:
                    comments.append(c.coordinate)
                if c.hyperlink:
                    hyperlinks.append(c.coordinate)
                if c.data_type == "f":
                    formulas.append(c.coordinate)
        num = pd.Series(valid_numeric)
        sheet = {
            "name": ws.title, "state": ws.sheet_state,
            "max_row_including_header": ws.max_row, "max_column": ws.max_column,
            "header": header, "data_rows": len(rows),
            "nonempty_data_rows": sum(any(v is not None for v in r) for r in rows),
            "all_blank_data_rows": sum(all(v is None for v in r) for r in rows),
            "exact_duplicate_rows_beyond_first": len(rows)-len(set(map(tuple, rows))),
            "cell_data_types_including_header": dict(data_types),
            "nulls_by_column": {h: sum(r[i] is None for r in rows) for i,h in enumerate(header)},
            "empty_strings_by_column": {h: sum(r[i] == "" for r in rows) for i,h in enumerate(header)},
            "timestamp": {
                "valid_datetime_count": len(valid_timestamps),
                "invalid_count": len(timestamps)-len(valid_timestamps),
                "timezone_aware_count": sum(t.tzinfo is not None for t in valid_timestamps),
                "first_in_sheet": iso(valid_timestamps[0]), "last_in_sheet": iso(valid_timestamps[-1]),
                "minimum": iso(min(valid_timestamps)), "maximum": iso(max(valid_timestamps)),
                "unique_count": len(set(valid_timestamps)),
                "duplicates_beyond_first": len(valid_timestamps)-len(set(valid_timestamps)),
                "strictly_increasing": all(b>a for a,b in zip(valid_timestamps,valid_timestamps[1:])),
                "adjacent_interval_seconds_counts": dict(deltas),
                "expected_naive_2024_quarter_hours": len(expected),
                "missing_from_expected_naive_grid_count": len(expected.difference(idx)),
                "extra_to_expected_naive_grid_count": len(idx.difference(expected)),
                "distinct_dates": len(per_day),
                "daily_row_count_distribution": dict(Counter(per_day.values())),
                "date_counts": dict(sorted(per_day.items())),
                "hypothetical_Europe_Brussels_localization": {
                    "not_a_timezone_inference": True,
                    "ambiguous_wall_times": [t.isoformat() for t in idx[ambiguous]],
                    "nonexistent_wall_times": [t.isoformat() for t in idx[nonexistent]],
                    "observed_rows_on_spring_transition": per_day["2024-03-31"],
                    "observed_rows_on_autumn_transition": per_day["2024-10-27"],
                },
            },
            "T_out": {
                "finite_numeric_count": len(valid_numeric),
                "nonnumeric_or_nonfinite_count": len(values)-len(valid_numeric),
                "distinct_values": len(set(valid_numeric)),
                "minimum": min(valid_numeric), "maximum": max(valid_numeric),
                "mean": num.mean(), "standard_deviation_sample": num.std(),
                "quantiles": {str(k):v for k,v in num.quantile([0,.01,.25,.5,.75,.99,1]).items()},
                "zero_count": sum(v == 0 for v in valid_numeric),
                "negative_count": sum(v < 0 for v in valid_numeric),
                "minimum_timestamp": iso(timestamps[values.index(min(valid_numeric))]),
                "maximum_timestamp": iso(timestamps[values.index(max(valid_numeric))]),
                "unit_not_inferred_by_this_audit": True,
            },
            "formatting_and_annotations": {
                "styles": dict(styles), "fills": dict(fills),
                "number_formats": dict(Counter(c.number_format for r in all_rows for c in r)),
                "comment_count": len(comments), "hyperlink_count": len(hyperlinks),
                "formula_count": len(formulas),
                "conditional_formatting_rule_groups": len(ws.conditional_formatting),
                "table_names": list(ws.tables), "merged_ranges": [str(r) for r in ws.merged_cells],
                "auto_filter": ws.auto_filter.ref,
                "hidden_rows": [i for i,d in ws.row_dimensions.items() if d.hidden],
                "hidden_columns": [i for i,d in ws.column_dimensions.items() if d.hidden],
            },
        }
        result["sheets"].append(sheet)
    wb.close()
    result["source_sha256_unchanged_after_read"] = digest_before == hashlib.sha256(path.read_bytes()).hexdigest()
    result["interpretation_limits"] = [
        "This dataset-specific audit compares naive timestamps with the quarter-hour grid for calendar 2024.",
        "Grid coverage and hypothetical Brussels localization do not establish UTC/local/standard-time semantics or the source's DST handling.",
        "The audit does not infer temperature units, station, provider, interval start/end, observation time, averaging, or interpolation.",
        "Descriptive metadata are checked for presence only; property values, saved locations, and annotation contents are omitted.",
        "Formatting and document metadata are not measurement provenance.",
        "This probe checks structure/completeness, not whether every value is an original measurement or accurate.",
    ]
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True, help="Input Excel workbook")
    parser.add_argument("--output", type=Path, required=True, help="Directory for summary.json")
    args = parser.parse_args()
    summary = inspect(args.workbook)
    args.output.mkdir(parents=True, exist_ok=True)
    destination = args.output / "summary.json"
    destination.write_text(json.dumps(summary, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(f"Wrote {destination}")
